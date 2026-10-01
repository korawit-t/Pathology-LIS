"""Spreadsheet rendering for list-style reports — the .xlsx counterpart to
pdf_service.

This exists because CSV cannot carry a cell *type*. An HN like "0012345" is
written to CSV correctly and then destroyed by Excel on open: the column is
parsed as numeric and the leading zeros are gone, silently turning a patient
identifier into a different one. The fix has to be a format that states the
type, so identifier columns here are written as text (`number_format = "@"`)
and Excel leaves them alone.
"""

import io
from typing import Any, Callable, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Excel's own ceiling is 1,048,576 rows; nothing here comes close, but a sheet
# name has a hard 31-character limit and invalid characters, so it is built
# rather than passed through.
_MAX_SHEET_TITLE = 31
# Excel rejects these in a sheet name: [ ] : * ? / \
_INVALID_SHEET_CHARS = str.maketrans({c: "-" for c in "[]:*?/" + chr(92)})

_HEADER_FILL = PatternFill("solid", fgColor="F0F0F0")
_HEADER_FONT = Font(bold=True)
_TITLE_FONT = Font(bold=True, size=14)
_LABEL_FONT = Font(bold=True)
_THIN = Side(style="thin", color="CCCCCC")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


class Column:
    """One output column.

    `as_text` forces Excel to treat the value as a string — set it for every
    identifier-shaped column (HN, accession no, CID), never for a real number
    you want to sort or sum.
    """

    def __init__(
        self,
        header: str,
        getter: Callable[[Any], Any],
        *,
        width: int = 18,
        as_text: bool = False,
        wrap: bool = False,
    ):
        self.header = header
        self.getter = getter
        self.width = width
        self.as_text = as_text
        self.wrap = wrap


def safe_sheet_title(title: str, default: str = "Report") -> str:
    cleaned = (title or "").translate(_INVALID_SHEET_CHARS).strip()
    return (cleaned or default)[:_MAX_SHEET_TITLE]


def build_table_workbook(
    *,
    title: str,
    criteria: Sequence[tuple[str, Any]],
    columns: Sequence[Column],
    rows: Sequence[Any],
    sheet_title: Optional[str] = None,
    footnotes: Sequence[str] = (),
) -> bytes:
    """Render a title, a label/value criteria block, a bordered table and
    optional footnotes into a single-sheet workbook, returned as bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = safe_sheet_title(sheet_title or title)

    span = max(len(columns), 2)
    r = 1

    ws.cell(row=r, column=1, value=title).font = _TITLE_FONT
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
    r += 2

    for label, value in criteria:
        ws.cell(row=r, column=1, value=label).font = _LABEL_FONT
        # Criteria values are free text (term lists, date ranges) and must not
        # be re-interpreted — "2026-01-01 ถึง 2026-03-31" is a label, not a date.
        cell = ws.cell(row=r, column=2, value="" if value is None else str(value))
        cell.number_format = "@"
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=span)
        r += 1

    r += 1
    header_row = r
    for i, col in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=i, value=col.header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = col.width
    r += 1

    for row_obj in rows:
        for i, col in enumerate(columns, start=1):
            raw = col.getter(row_obj)
            cell = ws.cell(row=r, column=i)
            if col.as_text:
                # The whole point of this module: write the string and say so,
                # so Excel does not re-parse "0012345" as 12345.
                cell.value = "" if raw is None else str(raw)
                cell.number_format = "@"
            else:
                cell.value = raw
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=col.wrap)
        r += 1

    # Freeze the header so a long case list stays readable while scrolling,
    # and let the reader filter the table without setting it up by hand.
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(columns))}{header_row + len(rows)}"
        )

    if footnotes:
        r += 1
        for note in footnotes:
            ws.cell(row=r, column=1, value=note).alignment = Alignment(wrap_text=True)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
            r += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
