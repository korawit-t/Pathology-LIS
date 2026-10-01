"""Router-level tests for app/routers/diagnosis_search.py. The matching logic
has full coverage in test_diagnosis_search.py — this is auth, the role gate,
and query-param wiring only."""

import uuid

from app.enums.surgical_diagnosis_enums import DiagnosisLevel
from app.models.surgical_diagnosis import SurgicalDiagnosis
from app.models.surgical_specimen import SurgicalSpecimen
from app.utils.time import local_now

from tests.factories import make_bare_case


def _reported_case(db, registrar_id, tag):
    case = make_bare_case(db, registrar_id=registrar_id)
    case.is_reported = True
    case.report_at = local_now()
    db.add(
        SurgicalSpecimen(
            case_id=case.id, specimen_label="A", specimen_name=f"Colon{tag}, biopsy{tag}"
        )
    )
    db.add(
        SurgicalDiagnosis(
            case_id=case.id,
            diagnosis_level=DiagnosisLevel.CASE,
            diagnosis=f"Adenocarcinoma{tag}",
            diagnosis_order=1,
            status="signed",
        )
    )
    db.commit()
    db.refresh(case)
    return case


class TestRbacAndAuth:
    def test_requires_authentication(self, client):
        r = client.get("/diagnosis-search", params={"specimen_terms": ["colon"]})
        assert r.status_code == 401

    def test_pathologist_is_allowed(self, pathologist_client):
        r = pathologist_client.get(
            "/diagnosis-search", params={"specimen_terms": ["colon"]}
        )
        assert r.status_code == 200

    def test_lab_manager_is_allowed(self, lab_manager_client):
        r = lab_manager_client.get(
            "/diagnosis-search", params={"specimen_terms": ["colon"]}
        )
        assert r.status_code == 200

    def test_clinician_is_rejected(self, clinician_client):
        r = clinician_client.get(
            "/diagnosis-search", params={"specimen_terms": ["colon"]}
        )
        assert r.status_code == 403


class TestValidation:
    def test_rejects_a_query_with_no_terms(self, pathologist_client):
        r = pathologist_client.get("/diagnosis-search")
        assert r.status_code == 400

    def test_rejects_a_query_whose_terms_are_all_blank(self, pathologist_client):
        r = pathologist_client.get(
            "/diagnosis-search", params={"specimen_terms": ["   ", ""]}
        )
        assert r.status_code == 400

    def test_rejects_an_unknown_match_mode(self, pathologist_client):
        r = pathologist_client.get(
            "/diagnosis-search",
            params={"specimen_terms": ["colon"], "match_mode": "sometimes"},
        )
        assert r.status_code == 422

    def test_rejects_an_unknown_date_field(self, pathologist_client):
        r = pathologist_client.get(
            "/diagnosis-search",
            params={"specimen_terms": ["colon"], "date_field": "collected"},
        )
        assert r.status_code == 422


class TestResponseWiring:
    def test_returns_the_matching_case_and_echoes_the_criteria(
        self, pathologist_client, db, admin_user
    ):
        registrar, _ = admin_user
        tag = uuid.uuid4().hex[:10]
        case = _reported_case(db, registrar.id, tag)

        r = pathologist_client.get(
            "/diagnosis-search",
            params={
                "specimen_terms": [f"colon{tag}", f"biopsy{tag}"],
                "diagnosis_terms": [f"adenocarcinoma{tag}"],
            },
        )
        body = r.json()

        assert r.status_code == 200
        assert body["total"] == 1
        assert body["items"][0]["accession_no"] == case.accession_no
        assert body["items"][0]["matched_specimens"] == [f"A: Colon{tag}, biopsy{tag}"]
        assert body["criteria"]["specimen_terms"] == [f"colon{tag}", f"biopsy{tag}"]
        assert body["criteria"]["match_mode"] == "all"
        assert body["criteria"]["date_field"] == "registered"

    def test_limit_caps_items_but_not_total(self, pathologist_client, db, admin_user):
        registrar, _ = admin_user
        tag = uuid.uuid4().hex[:10]
        for _ in range(2):
            _reported_case(db, registrar.id, tag)

        r = pathologist_client.get(
            "/diagnosis-search",
            params={"specimen_terms": [f"colon{tag}"], "limit": 1},
        )
        body = r.json()

        assert body["total"] == 2
        assert len(body["items"]) == 1


class TestXlsxExport:
    """The export exists in .xlsx rather than CSV for one reason: an HN like
    "0012345" is written to CSV correctly and then wrecked by Excel on open,
    which parses the column as numeric and drops the leading zeros — silently
    turning one patient's identifier into another's. These tests read the
    workbook back and assert the cell is text."""

    def _workbook(self, client, tag, **extra_params):
        import io

        from openpyxl import load_workbook

        params = {"specimen_terms": [f"colon{tag}"], **extra_params}
        r = client.get("/diagnosis-search/xlsx", params=params)
        assert r.status_code == 200, r.text
        assert "spreadsheetml.sheet" in r.headers["content-type"]
        assert ".xlsx" in r.headers["content-disposition"]
        return load_workbook(io.BytesIO(r.content))

    def _header_row(self, ws):
        for row in ws.iter_rows():
            if row[0].value == "ลำดับ":
                return row[0].row
        raise AssertionError("header row not found")

    def _cell(self, ws, header_row, header_text, offset=1):
        col = next(
            c.column for c in ws[header_row] if c.value == header_text
        )
        return ws.cell(row=header_row + offset, column=col)

    def test_hn_keeps_its_leading_zeros_as_a_text_cell(
        self, pathologist_client, db, admin_user
    ):
        registrar, _ = admin_user
        tag = uuid.uuid4().hex[:10]
        case = _reported_case(db, registrar.id, tag)
        case.hn = "0012345"
        db.commit()

        ws = self._workbook(pathologist_client, tag).active
        header_row = self._header_row(ws)
        hn = self._cell(ws, header_row, "HN")

        assert hn.value == "0012345"
        assert isinstance(hn.value, str)
        # "@" is Excel's text format — without it Excel re-parses on open.
        assert hn.number_format == "@"

    def test_accession_no_is_also_text(self, pathologist_client, db, admin_user):
        registrar, _ = admin_user
        tag = uuid.uuid4().hex[:10]
        case = _reported_case(db, registrar.id, tag)

        ws = self._workbook(pathologist_client, tag).active
        header_row = self._header_row(ws)
        acc = self._cell(ws, header_row, "Accession No.")

        assert acc.value == case.accession_no
        assert acc.number_format == "@"

    def test_carries_the_criteria_and_totals(self, pathologist_client, db, admin_user):
        registrar, _ = admin_user
        tag = uuid.uuid4().hex[:10]
        _reported_case(db, registrar.id, tag)

        ws = self._workbook(
            pathologist_client, tag, diagnosis_terms=[f"adenocarcinoma{tag}"]
        ).active
        labels = {
            row[0].value: row[1].value
            for row in ws.iter_rows(min_row=1, max_col=2)
            if row[0].value
        }

        assert labels["Specimen มีคำว่า"] == f"colon{tag}"
        assert labels["Diagnosis มีคำว่า"] == f"adenocarcinoma{tag}"
        assert labels["เงื่อนไขคำค้น"] == "ครบทุกคำ (AND)"
        assert labels["เคสที่เข้าเงื่อนไข"] == "1"

    def test_header_is_frozen_so_a_long_list_stays_readable(
        self, pathologist_client, db, admin_user
    ):
        registrar, _ = admin_user
        tag = uuid.uuid4().hex[:10]
        _reported_case(db, registrar.id, tag)

        ws = self._workbook(pathologist_client, tag).active

        assert ws.freeze_panes == f"A{self._header_row(ws) + 1}"

    def test_matched_specimen_and_diagnosis_collapse_into_one_cell_each(
        self, pathologist_client, db, admin_user
    ):
        registrar, _ = admin_user
        tag = uuid.uuid4().hex[:10]
        case = _reported_case(db, registrar.id, tag)
        db.add(
            SurgicalSpecimen(
                case_id=case.id, specimen_label="B", specimen_name=f"Colon{tag}, resection"
            )
        )
        db.commit()

        ws = self._workbook(pathologist_client, tag).active
        header_row = self._header_row(ws)
        specimens = self._cell(ws, header_row, "ชิ้นเนื้อที่ตรงเงื่อนไข").value

        assert f"A: Colon{tag}, biopsy{tag}" in specimens
        assert f"B: Colon{tag}, resection" in specimens
        assert "; " in specimens

    def test_rejects_a_query_with_no_terms(self, pathologist_client):
        assert pathologist_client.get("/diagnosis-search/xlsx").status_code == 400

    def test_clinician_is_rejected(self, clinician_client):
        r = clinician_client.get(
            "/diagnosis-search/xlsx", params={"specimen_terms": ["colon"]}
        )
        assert r.status_code == 403

    def test_requires_authentication(self, client):
        r = client.get("/diagnosis-search/xlsx", params={"specimen_terms": ["colon"]})
        assert r.status_code == 401
